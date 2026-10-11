# -*- coding: utf-8 -*-
from openpyxl import load_workbook
wb = load_workbook(r"D:\deepseekharness\datachange\kizuna\全站文案整理.xlsx")
print("Sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print("--- %s: %d rows x %d cols ---" % (name, ws.max_row, ws.max_column))
    hdr = [c.value for c in ws[1]]
    print("headers:", hdr)
