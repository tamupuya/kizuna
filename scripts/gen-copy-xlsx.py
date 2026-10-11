# -*- coding: utf-8 -*-
"""生成全站文案 Excel：界面文案 + 预设词库 两个 sheet"""
import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = r"D:\deepseekharness\datachange\kizuna"
data = json.load(open(ROOT + r"\scripts\copy-extract.json", encoding="utf-8"))

# 词库文件（整文件的 script string 视为预设词库，template 里的仍算界面文案）
DICT_FILES = {
    "InterestBuilder.vue": "兴趣分类/作品词库",
    "MbtiPicker.vue": "MBTI 词库",
    "ZodiacPicker.vue": "星座词库",
    "IdentityPicker.vue": "身份词库",
}

def is_dict(item, filename):
    base = filename.split("/")[-1]
    if base in DICT_FILES and item["zone"] == "script":
        return DICT_FILES[base]
    return None

# 意图预设 tags（IntentView script 里的预设数组）
def is_intent_preset(item, filename, text):
    if filename.split("/")[-1] == "IntentView.vue" and item["zone"] == "script":
        # 预设 tags 一般很短且是名词/身份词
        return True
    return False

def classify_big_small(item, text):
    """启发式判断大字/小字"""
    kind = item["kind"]
    if kind.startswith("attr:"):
        if kind == "attr:placeholder":
            return "占位符"
        if kind in ("attr:title", "attr:aria-label", "attr:alt", "attr:data-tip"):
            return "小字"
    t = text.strip()
    if kind == "text":
        # 纯变量插值模板不算
        if re.fullmatch(r"\{\{.*\}\}", t):
            return "动态"
        if len(t) <= 6 and not re.search(r"[，。！？：；]", t):
            return "大字"
        if re.search(r"请|提示|可|已|暂无|点击|如|建议|稍后|正在|需要|进行|输入|选择|上传|确认|取消|返回|设置|管理|查看|编辑|创建|发布|保存|退出|登录|注册", t):
            return "小字" if len(t) > 6 else "大字"
        if len(t) > 8:
            return "小字"
        return "大字"
    if kind == "string":
        return "小字"
    return "小字"

# 组装行
ui_rows = []
dict_rows = []
for f in data:
    fname = f["file"]
    for it in f["items"]:
        text = it["text"].strip()
        if not text:
            continue
        d = is_dict(it, fname)
        if d:
            dict_rows.append([fname, it["line"], d, text, ""])
        elif is_intent_preset(it, fname, text):
            dict_rows.append([fname, it["line"], "意图预设词库", text, ""])
        else:
            kind = it["kind"]
            kind_label = {
                "attr:placeholder": "占位符",
                "attr:title": "提示(title)",
                "attr:aria-label": "无障碍标签",
                "attr:alt": "图片替代文字",
                "attr:label": "标签(label)",
                "attr:data-tip": "提示",
            }.get(kind, "界面文字" if kind == "text" else "提示/文案")
            zone = "页面模板" if it["zone"] == "template" else "脚本逻辑"
            bs = classify_big_small(it, text)
            ui_rows.append([fname, it["line"], zone, kind_label, bs, text, ""])

# 按文件名排序
ui_rows.sort(key=lambda r: (r[0], r[1]))
dict_rows.sort(key=lambda r: (r[0], r[1]))

wb = Workbook()

# ---- Sheet 1: 界面文案 ----
ws = wb.active
ws.title = "界面文案"
headers = ["页面/组件", "行号", "区域", "类型", "大字/小字", "原文案", "你的修改"]
ws.append(headers)
for r in ui_rows:
    ws.append(r)

# ---- Sheet 2: 预设词库 ----
ws2 = wb.create_sheet("预设词库")
h2 = ["页面/组件", "行号", "词库分类", "原词条", "你的修改"]
ws2.append(h2)
for r in dict_rows:
    ws2.append(r)

# ---- 样式 ----
header_fill = PatternFill("solid", start_color="2B2A26")
header_font = Font(color="FFFFFF", bold=True, size=11)
thin = Side(style="thin", color="DDDDDD")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for sheet, ncols in ((ws, len(headers)), (ws2, len(h2))):
    for c in range(1, ncols + 1):
        cell = sheet.cell(row=1, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    # 冻结首行
    sheet.freeze_panes = "A2"
    # 列宽
    widths = {
        "A": 34, "B": 8, "C": 12, "D": 14, "E": 12, "F": 46, "G": 30,
    }
    for col, w in widths.items():
        if col <= get_column_letter(ncols):
            sheet.column_dimensions[col].width = w
    # 数据区边框 + 自动换行
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=ncols):
        for cell in row:
            cell.border = border
            if cell.column_letter in ("F", "G", "A"):
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    # 隔行浅色
    for i, row in enumerate(sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=ncols), start=2):
        if i % 2 == 0:
            for cell in row:
                cell.fill = PatternFill("solid", start_color="FAFAF8")

# Sheet2 列宽
ws2.column_dimensions["A"].width = 34
ws2.column_dimensions["B"].width = 8
ws2.column_dimensions["C"].width = 22
ws2.column_dimensions["D"].width = 40
ws2.column_dimensions["E"].width = 30

out = ROOT + r"\全站文案整理.xlsx"
wb.save(out)
print("界面文案:", len(ui_rows), "条")
print("预设词库:", len(dict_rows), "条")
print("已生成:", out)
