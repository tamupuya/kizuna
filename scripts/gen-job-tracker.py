# -*- coding: utf-8 -*-
"""生成求职追踪表.xlsx —— 求职闭环的配套工具（6 个 Sheet）"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\deepseekharness\datachange\kizuna\求职追踪表.xlsx"

# 样式
TITLE_FILL = PatternFill("solid", fgColor="2B2A26")
TITLE_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFF5DC")
HEAD_FILL = PatternFill("solid", fgColor="F4EEDE")
HEAD_FONT = Font(name="微软雅黑", size=10, bold=True, color="3C3C43")
BODY_FONT = Font(name="微软雅黑", size=10, color="3C3C43")
NOTE_FONT = Font(name="微软雅黑", size=10, color="8C8780")
THIN = Side(style="thin", color="D8D3C8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()

def style_header(ws, row=1, ncols=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER

def style_rows(ws, first=2, last=200, ncols=1):
    for r in range(first, last + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.alignment = WRAP
            cell.border = BORDER

def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

def title_row(ws, text, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.fill = TITLE_FILL
    c.font = TITLE_FONT
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

# ─────────────────────────────────────────
# Sheet 0: 使用说明
# ─────────────────────────────────────────
ws = wb.active
ws.title = "使用说明"
set_widths(ws, [22, 90])
rows = [
    ("求职闭环追踪表", "把求职当成做产品：获客（投递）→ 转化（邀约）→ 成交（Offer）→ 复盘 → 优化 → 再迭代。"),
    ("", ""),
    ("核心原则", "结果不可控，行为可控。只考核行为（投了多少、复盘了几次），不考核结果（拿没拿到 Offer）。"),
    ("", ""),
    ("Sheet：目标公司池", "用三维打分（匹配度/录取概率/想去程度）给每家打分，算出优先级，决定投递顺序。别海投，要精投。"),
    ("Sheet：投递记录", "每投一家立刻记一行。这是闭环的数据源——没有记录就没有迭代。"),
    ("Sheet：面试复盘", "每场面试后 30 分钟内填写。趁记忆新鲜，重点是最后一列「下次怎么改」。"),
    ("Sheet：问题库", "把每次被问的问题和你的最佳答案沉淀下来。同一问题答到第 5 次，你会明显变强。"),
    ("Sheet：漏斗看板", "每周统计一次。看哪一环转化率异常，对应用不同解法（详见下表）。"),
    ("", ""),
    ("漏斗诊断速查", "投递后已读不回 → 改打招呼语/换渠道"),
    ("", "已读但无邀约 → 简历按 JD 改关键词、量化、压到一页"),
    ("", "有邀约但一面挂 → 重练 3 分钟项目口述（录音自听）"),
    ("", "一面过终面挂 → 准备薪资话术、职业规划话术"),
    ("", "全程无人理 → 定位错了，下调公司标准/扩大行业范围"),
    ("", ""),
    ("每周节奏", "周日复盘 30min 定下周唯一改进项 ｜ 周一~五 每天投 3-5 家 + 记录 ｜ 每场面试后 30min 写复盘卡"),
]
for i, (a, b) in enumerate(rows, start=1):
    ws.cell(row=i, column=1, value=a).font = HEAD_FONT if a else BODY_FONT
    ws.cell(row=i, column=2, value=b).font = BODY_FONT if b else BODY_FONT
    ws.cell(row=i, column=2).alignment = WRAP

# ─────────────────────────────────────────
# Sheet 1: 目标公司池
# ─────────────────────────────────────────
ws = wb.create_sheet("目标公司池")
cols = ["优先级", "公司名", "岗位名称", "城市", "公司规模", "行业方向", "渠道",
        "匹配度(1-5)", "录取概率(1-5)", "想去程度(1-5)", "综合分", "JD链接", "备注"]
ws.append(cols)
title_row(ws, "目标公司池 —— 三维打分筛选（综合分 = 匹配×0.4 + 概率×0.35 + 想去×0.25）", len(cols))
ws.delete_rows(1)
ws.append(cols)
style_header(ws, 1, len(cols))
samples = [
    ["", "（示例）某ACG社区公司", "产品经理（初级）", "上海", "50-200人", "泛二次元社区", "BOSS直聘", 5, 4, 5, None, "", "JD要求「有完整项目经验优先」→ 作品集可打"],
]
for s in samples:
    ws.append(s)
# 公式列 K
for r in range(2, 202):
    ws.cell(row=r, column=11, value=f"=IF(COUNT(H{r}:J{r})=3,H{r}*0.4+I{r}*0.35+J{r}*0.25,\"\")")
style_rows(ws, 2, 201, len(cols))
for r in range(2, 202):
    for c in (1, 4, 5, 6, 7, 8, 9, 10, 11):
        ws.cell(row=r, column=c).alignment = CENTER
    ws.cell(row=r, column=11).font = Font(name="微软雅黑", size=10, bold=True, color="C44545")
set_widths(ws, [8, 22, 20, 8, 12, 16, 12, 11, 12, 11, 9, 24, 34])
ws.freeze_panes = "A2"

# ─────────────────────────────────────────
# Sheet 2: 投递记录
# ─────────────────────────────────────────
ws = wb.create_sheet("投递记录")
cols = ["投递日期", "公司名", "岗位名称", "渠道", "简历版本", "打招呼语版本",
        "状态", "最近更新", "邀约日期", "面试轮次", "结果", "未通过原因/卡在哪", "下一步动作"]
ws.append(cols)
style_header(ws, 1, len(cols))
ws.append(["2026-09-22", "（示例）某ACG社区公司", "产品经理（初级）", "BOSS直聘", "PM版", "v1",
           "已投递", "2026-09-22", "", "", "", "", "3天后跟进"])
style_rows(ws, 2, 201, len(cols))
for r in range(2, 202):
    for c in (1, 4, 7, 8, 9, 10, 11):
        ws.cell(row=r, column=c).alignment = CENTER
# 状态下拉
from openpyxl.worksheet.datavalidation import DataValidation
dv = DataValidation(type="list",
                    formula1='"目标池,已投递,已读未回,已邀约,一面,二面,终面,已Offer,已拒绝,已放弃"',
                    allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"G2:G201")
set_widths(ws, [12, 22, 20, 12, 11, 14, 11, 12, 12, 10, 10, 30, 26])
ws.freeze_panes = "A2"

# ─────────────────────────────────────────
# Sheet 3: 面试复盘
# ─────────────────────────────────────────
ws = wb.create_sheet("面试复盘")
cols = ["面试日期", "公司", "岗位", "轮次", "面试形式", "被问的问题（原话）",
        "我答得好的", "我卡住的", "面试官反应/信号", "下次具体怎么改", "结果"]
ws.append(cols)
style_header(ws, 1, len(cols))
ws.append(["2026-09-25", "（示例）某ACG社区公司", "产品经理（初级）", "一面", "线上视频",
           "「你的项目没上线，怎么验证设计有效？」",
           "隐私四层讲得清楚",
           "被追问数据时有点慌",
           "面试官点头，但追问了两次数据",
           "补充「三层验证」框架：外部对标验证 + 用户直觉 + 逻辑自洽",
           "待定"])
style_rows(ws, 2, 201, len(cols))
for r in range(2, 202):
    for c in (1, 4, 5, 11):
        ws.cell(row=r, column=c).alignment = CENTER
set_widths(ws, [12, 20, 18, 8, 10, 40, 22, 22, 24, 34, 10])
ws.freeze_panes = "A2"

# ─────────────────────────────────────────
# Sheet 4: 问题库
# ─────────────────────────────────────────
ws = wb.create_sheet("问题库")
cols = ["分类", "问题", "我的最佳答案（持续迭代）", "答过几次", "最近更新", "状态"]
ws.append(cols)
style_header(ws, 1, len(cols))
seeds = [
    ["通用", "请做一下自我介绍", "30秒版 + 3分钟版（见 docs/求职闭环方法论.md §3.2）", 0, "", "待练"],
    ["通用", "为什么想做产品经理？", "做项目过程中发现自己更享受「决定做什么」而非「怎么实现」", 0, "", "待练"],
    ["项目", "你的项目没上线，怎么验证设计有效？", "三层验证：外部对标验证 + 用户直觉 + 逻辑自洽；埋点已设计好", 0, "", "待练"],
    ["项目", "代码是你写的吗？", "AI辅助实现，产品决策/需求/验收由我主导；2026年AI PM要的是懂产品+会用AI", 0, "", "待练"],
    ["项目", "为什么砍掉社交动态功能？", "第一性原理：产品定位是「连接」不是「社交」，feed会稀释核心价值", 0, "", "待练"],
    ["项目", "对标产品上线了和你相似的设计，你怎么看？", "方向被商业化验证，说明判断一致；强调时间线，不说抄袭", 0, "", "待练"],
    ["公司", "你怎么看我们公司的产品？", "准备2-3个「如果我做会怎么改」的建议", 0, "", "待练"],
    ["薪资", "你的期望薪资是多少？", "先问岗位薪资区间，再给区间（不先报死价）", 0, "", "待练"],
]
for s in seeds:
    ws.append(s)
style_rows(ws, 2, 201, len(cols))
for r in range(2, 202):
    for c in (1, 4, 5, 6):
        ws.cell(row=r, column=c).alignment = CENTER
set_widths(ws, [10, 34, 56, 10, 12, 10])
ws.freeze_panes = "A2"

# ─────────────────────────────────────────
# Sheet 5: 漏斗看板
# ─────────────────────────────────────────
ws = wb.create_sheet("漏斗看板")
cols = ["周次", "本周投递", "累计投递", "邀约数", "简历通过率", "实面数", "到面率",
        "进二面数", "一面通过率", "Offer数", "本周最大收获", "本周最大卡点", "下周唯一改进项"]
ws.append(cols)
style_header(ws, 1, len(cols))
# 预置周次 + 公式
for i, wk in enumerate(["第1周", "第2周", "第3周", "第4周", "第5周", "第6周", "第7周", "第8周"], start=2):
    ws.cell(row=i, column=1, value=wk).alignment = CENTER
    ws.cell(row=i, column=3, value=f"=SUM($B$2:B{i})")
    ws.cell(row=i, column=5, value=f"=IF(C{i}=0,\"\",D{i}/C{i})")
    ws.cell(row=i, column=7, value=f"=IF(D{i}=0,\"\",F{i}/D{i})")
    ws.cell(row=i, column=9, value=f"=IF(F{i}=0,\"\",H{i}/F{i})")
style_rows(ws, 2, 60, len(cols))
for r in range(2, 61):
    for c in (1, 2, 3, 4, 5, 6, 7, 8, 9, 10):
        ws.cell(row=r, column=c).alignment = CENTER
    for c in (5, 7, 9):
        ws.cell(row=r, column=c).number_format = "0%"
# 参考转化率提示
ws.cell(row=62, column=1, value="参考区间").font = HEAD_FONT
ws.cell(row=62, column=1).fill = HEAD_FILL
ws.cell(row=63, column=1, value="简历通过率").font = BODY_FONT
ws.cell(row=63, column=2, value="10%~20%").font = BODY_FONT
ws.cell(row=64, column=1, value="到面率").font = BODY_FONT
ws.cell(row=64, column=2, value="60%~80%").font = BODY_FONT
ws.cell(row=65, column=1, value="一面通过率").font = BODY_FONT
ws.cell(row=65, column=2, value="30%~40%").font = BODY_FONT
ws.cell(row=66, column=1, value="说明").font = NOTE_FONT
ws.cell(row=66, column=2, value="参考值仅用于判断哪一环异常；你自己的历史数据才是真基准").font = NOTE_FONT
set_widths(ws, [10, 10, 10, 9, 12, 9, 9, 10, 12, 9, 26, 26, 30])
ws.freeze_panes = "A2"

wb.save(OUT)
print("OK ->", OUT, os.path.getsize(OUT), "bytes")
print("Sheets:", wb.sheetnames)
